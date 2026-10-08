"""抓台股行情，寫進 market.json，給儀表板上方四格使用。

資料來源：
  - 加權指數、個股即時價：證交所「基本市況報導」mis.twse.com.tw（上市 tse_、上櫃 otc_）
  - 大盤成交金額：證交所 MI_5MINS（每 5 秒委託成交統計）
  - 證交所連不上時，指數和個股改用 Yahoo Finance 備援
追蹤清單：讀 data.xlsx 的「標的清單」（市場＝台股），名稱取「公司檔案」。
"""
import json, ssl, time, urllib.request, http.cookiejar
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TW = timezone(timedelta(hours=8))
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36"
jar = http.cookiejar.CookieJar()
opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))


def get_json(url, referer=None, timeout=15):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "application/json", **({"Referer": referer} if referer else {})})
    with opener.open(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))


def num(v):
    try:
        f = float(str(v).replace(",", ""))
        return f if f > 0 else None
    except (TypeError, ValueError):
        return None


def watchlist():
    wb = openpyxl.load_workbook(ROOT / "data.xlsx", data_only=True, read_only=True)
    names = {}
    if "公司檔案" in wb.sheetnames:
        rows = list(wb["公司檔案"].iter_rows(values_only=True))
        h = list(rows[0])
        for r in rows[1:]:
            if r[h.index("代號")]:
                names[str(r[h.index("代號")]).strip().upper()] = r[h.index("名稱")]
    rows = list(wb["標的清單"].iter_rows(values_only=True))
    h = list(rows[1])  # 第 2 列是欄位標題
    codes = []
    for r in rows[2:]:
        code, mkt = r[h.index("代號")], r[h.index("市場")]
        cat = r[h.index("分類")] if "分類" in h else None
        if not code or cat == "已出場" or (mkt and mkt != "台股"):
            continue
        c = str(code).strip()
        if c.endswith(".0"):
            c = c[:-2]
        if c not in codes:
            codes.append(c)
    return codes, names


def from_mis(codes):
    """回傳 ({'t00': (價, 昨收)}, {代號: (價, 昨收, 名稱)})"""
    try:
        opener.open(urllib.request.Request("https://mis.twse.com.tw/stock/index.jsp", headers={"User-Agent": UA}), timeout=15).read()
    except Exception:
        pass
    chans = ["tse_t00.tw"] + [f"{ex}_{c}.tw" for c in codes for ex in ("tse", "otc")]
    idx, stocks = {}, {}
    for i in range(0, len(chans), 40):
        url = ("https://mis.twse.com.tw/stock/api/getStockInfo.jsp?json=1&delay=0&_=%d&ex_ch=" % int(time.time() * 1000)) + "|".join(chans[i:i + 40])
        data = get_json(url, referer="https://mis.twse.com.tw/stock/index.jsp")
        for m in data.get("msgArray", []):
            price = num(m.get("z"))
            if price is None:  # 剛好沒成交時，用最佳買價代替
                price = num((m.get("b") or "").split("_")[0])
            prev = num(m.get("y"))
            if m.get("c") == "t00":
                idx["t00"] = (price, prev)
            elif m.get("c") in codes:
                stocks[m["c"]] = (price, prev, m.get("n"))
        time.sleep(1)
    return idx, stocks


def yahoo(symbol):
    d = get_json(f"https://query1.finance.yahoo.com/v8/finance/chart/{symbol}?range=1d&interval=1d")
    meta = d["chart"]["result"][0]["meta"]
    return num(meta.get("regularMarketPrice")), num(meta.get("chartPreviousClose") or meta.get("previousClose"))


def turnover():
    d = get_json("https://www.twse.com.tw/exchangeReport/MI_5MINS?response=json")
    rows = d.get("data") or []
    if not rows:
        return None
    last = rows[-1]
    f = d.get("fields") or []
    amt = num(last[f.index("累積成交金額")]) if "累積成交金額" in f else num(last[-1])  # 單位：百萬元
    return {"amount_e8": round(amt / 100, 1), "time": last[0]} if amt else None


def main():
    codes, names = watchlist()
    idx, stocks = {}, {}
    try:
        idx, stocks = from_mis(codes)
    except Exception as e:
        print("證交所即時資料失敗，改用 Yahoo：", e)
    if not idx.get("t00") or not idx["t00"][0]:
        try:
            idx["t00"] = yahoo("%5ETWII")
        except Exception as e:
            print("指數備援失敗：", e)
    for c in codes:
        if c in stocks and stocks[c][0] and stocks[c][1]:
            continue
        for suffix in (".TW", ".TWO"):
            try:
                p, y = yahoo(c + suffix)
                if p and y:
                    stocks[c] = (p, y, None)
                    break
            except Exception:
                pass

    out = {"updated": datetime.now(TW).isoformat(timespec="seconds")}
    p, y = idx.get("t00", (None, None))
    if p and y:
        out["index"] = {"value": round(p, 2), "change": round(p - y, 2), "pct": round(p / y - 1, 5)}
    movers = []
    for c, (price, prev, n) in stocks.items():
        if price and prev:
            movers.append({"code": c, "name": names.get(c.upper()) or n or c, "price": price, "pct": round(price / prev - 1, 5)})
    if movers:
        movers.sort(key=lambda m: m["pct"])
        out["top"], out["bottom"] = movers[-1], movers[0]
        out["count"] = len(movers)
    try:
        t = turnover()
        if t:
            out["turnover"] = {"amount_e8": t["amount_e8"]}
    except Exception as e:
        print("成交金額失敗：", e)
    out["time_label"] = datetime.now(TW).strftime("%m/%d %H:%M")

    path = ROOT / "market.json"
    old = {}
    if path.exists():
        try:
            old = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            pass
    strip = lambda d: {k: v for k, v in d.items() if k not in ("updated", "time_label")}
    if "index" not in out and old.get("index"):
        print("這次沒抓到資料，保留上一次的行情")
        return
    if strip(out) == strip(old):
        print("行情沒有變動")
        return
    path.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(out, ensure_ascii=False))


if __name__ == "__main__":
    main()
